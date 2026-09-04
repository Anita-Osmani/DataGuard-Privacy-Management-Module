from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn

OUT = Path('outputs/Kapitulli_1_Hyrje_DataGuard.docx')

doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)
section.left_margin = Cm(3)
section.right_margin = Cm(3)
section.header_distance = Cm(1.25)
section.footer_distance = Cm(1.25)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Times New Roman'
normal._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
normal._element.rPr.rFonts.set(qn('w:cs'), 'Times New Roman')
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.first_line_indent = Cm(1.25)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for name, size in [('Heading 1', 14), ('Heading 2', 12)]:
    style = styles[name]
    style.font.name = 'Times New Roman'
    style._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
    style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = None
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.first_line_indent = Cm(0)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

def add_heading(text, level):
    doc.add_paragraph(text, style=f'Heading {level}')

def add_para(text):
    p = doc.add_paragraph(style='Normal')
    p.add_run(text)
    return p

add_heading('1. HYRJE', 1)

add_heading('1.1 Konteksti i temës', 2)
add_para('Zhvillimi i shpejtë i shërbimeve digjitale ka bërë që organizatat, përfshirë ndërmarrjet e vogla dhe të mesme (NVM), të mbledhin dhe të përdorin të dhëna personale në veprimtarinë e përditshme. Edhe një biznes i vogël mund të ruajë të dhëna të klientëve, si emri, adresa elektronike, numri i telefonit dhe adresa, për qëllime të komunikimit, ofrimit të shërbimit ose administrimit të marrëdhënies me klientin. Përdorimi i këtyre të dhënave krijon nevojën që informacioni të menaxhohet me përgjegjësi, në mënyrë të kuptueshme për përdoruesin dhe të kontrollueshme për administratorin e sistemit.')
add_para('Privatësia nuk duhet të shihet vetëm si detyrim formal apo si tekst i vendosur në fund të një faqeje interneti. Ajo lidhet me mënyrën se si përdoruesi informohet, jep miratimin, tërheq miratimin, kërkon kopjen e të dhënave të veta dhe kërkon fshirjen e tyre. Po aq e rëndësishme është që organizata të mund të tregojë se cilat veprime janë kryer ndaj të dhënave dhe nga cili përdorues. Kjo kërkon mekanizma të thjeshtë për menaxhimin e miratimeve, dokumentimin e politikave, regjistrimin e aktiviteteve dhe administrimin e periudhave të ruajtjes së të dhënave.')
add_para('Rregullorja e Përgjithshme e Bashkimit Evropian për Mbrojtjen e të Dhënave (GDPR) paraqet një kornizë të rëndësishme referuese për mbrojtjen e të dhënave personale. Ajo thekson ligjshmërinë, transparencën, kufizimin e qëllimit, minimizimin e të dhënave dhe kufizimin e periudhës së ruajtjes [1]. Edhe pse zbatimi praktik ndryshon sipas kontekstit juridik dhe organizativ, këto parime janë të dobishme për projektimin e aplikacioneve që respektojnë privatësinë dhe krijojnë besim te përdoruesit.')
add_para('Në praktikë, NVM-të zakonisht nuk kanë burime për sisteme të mëdha të menaxhimit të privatësisë ose për procese të ndërlikuara administrative. Për këtë arsye, ekziston nevoja për një modul të qartë, me funksione të kuptueshme dhe të integrueshme në aplikacione ekzistuese. Ky punim trajton projektimin dhe zhvillimin e një zgjidhjeje të tillë, të quajtur DataGuard, e cila demonstron mënyrën se si funksionet bazë të privatësisë mund të vendosen në një aplikacion për menaxhimin e klientëve.')

add_heading('1.2 Përkufizimi i problemit', 2)
add_para('Problemi kryesor i trajtuar në këtë punim është mungesa e një mekanizmi të centralizuar dhe të thjeshtë për menaxhimin e privatësisë në aplikacionet e NVM-ve. Në shumë raste, të dhënat e klientëve ruhen në formularë, tabela ose aplikacione të brendshme pa një lidhje të qartë ndërmjet politikës së privatësisë, miratimit të përdoruesit dhe veprimeve të kryera mbi të dhënat. Si pasojë, organizata mund të ketë vështirësi të përgjigjet në pyetje praktike: Cilin version të politikës e ka pranuar përdoruesi? A e ka tërhequr miratimin? Kush e ka ndryshuar një të dhënë? Sa gjatë duhet të ruhet një kategori e caktuar e të dhënave?')
add_para('Mungesa e këtyre mekanizmave nuk do të thotë domosdoshmërisht se çdo biznes vepron në mënyrë të papërgjegjshme. Shpesh problemi lidhet me kufizimet e kohës, buxhetit dhe njohurive teknike. Megjithatë, proceset e paformalizuara e bëjnë më të vështirë transparencën, mbikëqyrjen dhe zbatimin e kërkesave të përdoruesve. Për shembull, ruajtja e një politike pa histori versionsh nuk e lejon organizatën të dëshmojë se çfarë teksti ka qenë në fuqi në momentin kur është dhënë miratimi.')
add_para('Një aspekt tjetër i problemit është mungesa e gjurmueshmërisë së veprimeve. Veprime të tilla si hyrja në sistem, shikimi i një të dhëne, krijimi ose ndryshimi i një politike, regjistrimi i miratimit, eksportimi i të dhënave dhe fshirja e një klienti janë relevante për administrimin e privatësisë. Pa audit logs, administratori nuk ka një pasqyrë të përmbledhur mbi këto aktivitete. Prandaj, zgjidhja e propozuar synon të ofrojë një regjistër të thjeshtë, të filtrueshëm dhe të kuptueshëm të veprimeve kryesore.')
add_para('Problemi përfshin gjithashtu ruajtjen dhe përpunimin e kërkesave të përdoruesve. Përdoruesi duhet të ketë mundësi të kërkojë kopjen e të dhënave të veta ose fshirjen e tyre, ndërsa administratori duhet të mund t’i shqyrtojë dhe t’i përpunojë këto kërkesa. GDPR njeh të drejtën e qasjes në të dhënat personale dhe të drejtën për fshirje në rrethana të caktuara [1]. Në kuadër të këtij prototipi, këto të drejta paraqiten përmes rrjedhave të thjeshta dhe të kontrollueshme, pa pretenduar automatizim të plotë juridik ose operacional.')

add_heading('1.3 Qëllimi dhe objektivat e punimit', 2)
add_para('Qëllimi i këtij punimi është të projektojë dhe të zhvillojë një modul funksional për menaxhimin e miratimeve të privatësisë, regjistrimin e qasjeve dhe ruajtjen e të dhënave për aplikacionet e NVM-ve. Moduli synon të jetë i thjeshtë për përdorim, i qartë në ndërfaqe dhe i mjaftueshëm për të demonstruar funksionet kryesore që kërkohen nga një sistem bazë i menaxhimit të privatësisë.')
add_para('Objektivi i parë është krijimi i mekanizmit për politikat e privatësisë dhe versionimin e tyre. Administratori mund të krijojë një politikë, të përcaktojë datën e hyrjes në fuqi, të shikojë përmbajtjen e saj dhe të krijojë versione të reja duke ruajtur historikun e mëparshëm. Objektivi i dytë është menaxhimi i miratimeve, ku përdoruesi mund ta pranojë ose ta tërheqë miratimin dhe sistemi ruan versionin konkret të politikës me të cilin lidhet vendimi.')
add_para('Objektivi i tretë është krijimi i një mekanizmi të thjeshtë auditimi. Në këtë mekanizëm regjistrohen veprimet relevante me datë, përdorues, veprim, burim dhe status. Objektivi i katërt është përcaktimi i politikave të ruajtjes së të dhënave sipas kategorisë së të dhënës, periudhës së ruajtjes dhe veprimit të planifikuar, si fshirje, arkivim ose ruajtje. Objektivi i pestë është implementimi i rrjedhave për kërkesë dhe eksportim të të dhënave, si dhe kërkesë për fshirje.')
add_para('Objektivi përfundimtar është demonstrimi i integrimit të modulit në një aplikacion të thjeshtë SME për menaxhimin e klientëve. Kjo pjesë vërteton se funksionet e privatësisë nuk janë të shkëputura nga aplikacioni kryesor, por mund të mbështesin menaxhimin e të dhënave të klientëve në përdorimin e përditshëm.')

add_heading('1.4 Pyetjet kërkimore', 2)
add_para('Punimi udhëhiqet nga pyetja kryesore: si mund të ndërtohet një modul i thjeshtë dhe i përdorshëm që ndihmon aplikacionet e NVM-ve të menaxhojnë miratimet e privatësisë, qasjet në të dhëna dhe periudhat e ruajtjes së tyre? Kjo pyetje ndahet në disa pyetje mbështetëse: si ruhet lidhja ndërmjet përdoruesit, politikës dhe versionit të pranuar; cilat veprime duhet të regjistrohen në audit log; si mund të paraqiten politikat e ruajtjes në një formë të kuptueshme; dhe si mund të realizohen kërkesat për eksportim ose fshirje pa e komplikuar sistemin.')
add_para('Përgjigjja ndaj këtyre pyetjeve ndërtohet përmes analizës së kërkesave, projektimit të arkitekturës së sistemit, krijimit të modelit të të dhënave dhe testimit të rrjedhave kryesore. Qasja e zgjedhur nuk synon të zëvendësojë një platformë të madhe për pajtueshmëri ligjore, por të tregojë një bazë të zbatueshme dhe të zgjerueshme për një biznes të vogël ose të mesëm.')

add_heading('1.5 Fushëveprimi i sistemit', 2)
add_para('Fushëveprimi i zgjidhjes përfshin dy role: administratorin dhe përdoruesin. Administratori ka qasje në panelin e menaxhimit, ku administron politikat e privatësisë, versionet, miratimet, audit logs, politikat e ruajtjes, kërkesat e të dhënave dhe klientët e aplikacionit demonstrues. Përdoruesi ka qasje në hapësirën “My Privacy”, ku mund të shikojë politikën aktuale, të japë ose të tërheqë miratimin dhe të dërgojë kërkesë për të dhënat e veta.')
add_para('Sistemi mbulon ruajtjen e profileve të përdoruesve, politikave të privatësisë, versioneve të politikave, miratimeve, regjistrimeve të auditimit, politikave të ruajtjes, kërkesave për të dhëna dhe klientëve. Zgjidhja përdor React, TypeScript, Vite dhe Tailwind CSS në ndërfaqen e përdoruesit, ndërsa Supabase dhe PostgreSQL përdoren për autentifikim, ruajtjen e të dhënave dhe qasjen në API. Kontrolli bazë i qasjes bazohet në rolet Admin dhe User.')
add_para('Në mënyrë të qëllimshme, fushëveprimi nuk përfshin funksione të avancuara si zbulimi automatik i anomalive, blockchain, enkriptimi i veçantë i menaxhuar nga aplikacioni, mikroshërbimet, multi-tenancy ose punë të automatizuara në sfond për fshirje. Procesimi i politikave të ruajtjes është manual në prototip. Po ashtu, kërkesa për fshirje regjistrohet dhe menaxhohet nga administratori; anonimizimi i automatizuar nuk implementohet dhe mbetet mundësi për zgjerim të ardhshëm.')

add_heading('1.6 Metodologjia e zhvillimit', 2)
add_para('Zhvillimi i sistemit është realizuar me një qasje iterative. Fillimisht janë identifikuar kërkesat funksionale dhe aktorët kryesorë të sistemit. Më pas është projektuar modeli i bazës së të dhënave duke përcaktuar entitetet dhe marrëdhëniet kryesore, si lidhja ndërmjet politikës së privatësisë, versionit të saj dhe miratimit të përdoruesit. Në fazën pasuese është ndërtuar ndërfaqja e përdoruesit dhe janë implementuar rrjedhat e veprimeve për secilin modul.')
add_para('Për realizimin teknik është përdorur React me TypeScript për krijimin e komponentëve të ripërdorshëm dhe për menaxhimin e gjendjes së ndërfaqes. Vite është përdorur për mjedisin e zhvillimit, ndërsa Tailwind CSS për stilizimin e ndërfaqes. Supabase shërben si platformë backend, duke ofruar Supabase Auth për autentifikim, PostgreSQL për ruajtjen e të dhënave dhe API për komunikimin e aplikacionit me bazën e të dhënave. Rregullat e Row Level Security përdoren për të kufizuar qasjen bazuar në rolin e përdoruesit.')
add_para('Pas implementimit, funksionet janë verifikuar përmes skenarëve të përdorimit. Këta skenarë përfshijnë krijimin e politikës, shtimin e një versioni të ri, pranimin dhe tërheqjen e miratimit, shtimin dhe ndryshimin e klientit, krijimin e kërkesës për të dhëna, përpunimin e saj nga administratori dhe kontrollin e regjistrimeve në audit log. Rezultatet e testimit paraqiten më hollësisht në kapitullin e nëntë.')

add_heading('1.7 Struktura e punimit', 2)
add_para('Punimi është organizuar në dhjetë kapituj. Kapitulli i parë paraqet kontekstin, problemin, qëllimin, objektivat, pyetjet kërkimore, fushëveprimin dhe metodologjinë e zhvillimit. Kapitulli i dytë trajton kontekstin ligjor dhe rishikimin e literaturës mbi privatësinë, mbrojtjen e të dhënave, miratimin dhe auditimin. Kapitulli i tretë analizon kërkesat funksionale dhe jofunksionale të sistemit, aktorët dhe rastet e përdorimit.')
add_para('Kapitulli i katërt shpjegon arkitekturën e zgjidhjes, teknologjitë e përdorura, modelin e bazës së të dhënave, ndërfaqen dhe kontrollin e qasjes. Kapitujt e pestë, të gjashtë dhe të shtatë paraqesin në detaje modulin e miratimeve dhe politikave të privatësisë, modulin e auditimit, si dhe politikat e ruajtjes, eksportimit dhe fshirjes së të dhënave. Kapitulli i tetë përshkruan integrimin në aplikacionin demonstrues SME për menaxhimin e klientëve.')
add_para('Në kapitullin e nëntë paraqiten testimi funksional, kontrolli i autentifikimit dhe autorizimit, konsideratat e sigurisë dhe vlerësimi i rezultateve. Në fund, kapitulli i dhjetë përmbledh arritjet e punimit, kufizimet e prototipit dhe mundësitë për punë të ardhshme. Kjo strukturë mundëson që pjesa teorike, analitike dhe praktike të lidhen në mënyrë të qartë me objektivin e përgjithshëm të punimit.')

add_heading('Referenca të përdorura në këtë kapitull', 2)
refs = [
    '[1] European Parliament and Council of the European Union. Regulation (EU) 2016/679 (General Data Protection Regulation), 27 April 2016. EUR-Lex. Available: https://eur-lex.europa.eu/eli/reg/2016/679/oj',
    '[2] European Data Protection Board. Guidelines 05/2020 on consent under Regulation 2016/679, Version 1.1, 4 May 2020. Available: https://www.edpb.europa.eu/our-work-tools/our-documents/guidelines/guidelines-052020-consent-under-regulation-2016679_en'
]
for ref in refs:
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0)
    p.paragraph_format.space_after = Pt(3)
    p.add_run(ref)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run('DataGuard - Punim diplome')
fr.font.name = 'Times New Roman'
fr._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
fr.font.size = Pt(10)

OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(OUT)
