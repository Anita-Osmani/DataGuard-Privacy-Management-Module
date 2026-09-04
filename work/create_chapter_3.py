from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path('outputs/Kapitulli_3_Analiza_e_Kerkesave.docx')
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.5)
sec.left_margin, sec.right_margin = Cm(3), Cm(3)
sec.header_distance, sec.footer_distance = Cm(1.25), Cm(1.25)

normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
normal._element.rPr.rFonts.set(qn('w:cs'), 'Times New Roman')
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.first_line_indent = Cm(1.25)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for name, size in [('Heading 1', 14), ('Heading 2', 12)]:
    s = doc.styles[name]
    s.font.name = 'Times New Roman'
    s._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
    s._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    s.font.size, s.font.bold = Pt(size), True
    s.paragraph_format.line_spacing = 1.5
    s.paragraph_format.space_before, s.paragraph_format.space_after = Pt(12), Pt(6)
    s.paragraph_format.first_line_indent = Cm(0)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

def heading(text, level): doc.add_paragraph(text, style=f'Heading {level}')
def para(text):
    p = doc.add_paragraph(style='Normal'); p.add_run(text); return p
def run_font(run, bold=False, size=10):
    run.font.name = 'Times New Roman'; run._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman'); run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman'); run.font.size = Pt(size); run.font.bold = bold
def set_cell_text(cell, text, bold=False):
    cell.text = ''
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.15; p.paragraph_format.first_line_indent = Cm(0); p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text); run_font(run, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT; table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in('w:tblW')
    tbl_w.set(qn('w:w'), str(sum(widths))); tbl_w.set(qn('w:type'), 'dxa')
    ind = tbl_pr.first_child_found_in('w:tblInd')
    if ind is None:
        ind = OxmlElement('w:tblInd'); tbl_pr.append(ind)
    ind.set(qn('w:w'), '120'); ind.set(qn('w:type'), 'dxa')
    for grid_col, width in zip(table._tbl.tblGrid.gridCol_lst, widths): grid_col.set(qn('w:w'), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_w = cell._tc.tcPr.tcW
            tc_w.set(qn('w:w'), str(width)); tc_w.set(qn('w:type'), 'dxa')
    tr_pr = table.rows[0]._tr.get_or_add_trPr(); hdr = OxmlElement('w:tblHeader'); hdr.set(qn('w:val'), 'true'); tr_pr.append(hdr)

heading('3. ANALIZA E KËRKESAVE TË SISTEMIT', 1)

heading('3.1 Përshkrimi i sistemit', 2)
para('DataGuard është një modul i privatësisë për aplikacionet e ndërmarrjeve të vogla dhe të mesme. Qëllimi i tij është të mbështesë menaxhimin e të dhënave personale pa kërkuar një infrastrukturë të ndërlikuar ose procese enterprise. Moduli integrohet me një aplikacion demonstrues për menaxhimin e klientëve dhe krijon një vend të centralizuar për politikat e privatësisë, miratimet, auditimin, politikat e ruajtjes dhe kërkesat e përdoruesve për të dhënat e veta.')
para('Sistemi është projektuar si aplikacion web me ndërfaqe të përgjegjshme për ekrane të ndryshme. Përdoruesit identifikohen nëpërmjet Supabase Auth dhe marrin qasje sipas rolit të tyre. Roli Administrator përdor panelin e menaxhimit, ndërsa roli User përdor hapësirën “My Privacy”. Kjo ndarje synon të shmangë paraqitjen e funksioneve administrative për përdoruesin e zakonshëm dhe të bëjë rrjedhat e përdorimit më të thjeshta.')
para('Në nivel të të dhënave, sistemi ruan profilet e përdoruesve, politikat e privatësisë, versionet e tyre, miratimet, audit logs, politikat e ruajtjes, kërkesat për të dhëna dhe klientët e aplikacionit demonstrues. Çdo funksion i ndërfaqes lidhet me njërin nga këto entitete. Për shembull, pranimi i politikës krijon regjistrim të miratimit të lidhur me versionin konkret të politikës, ndërsa kërkesa për fshirje krijon regjistrim me status fillestar Pending.')
para('Sistemi nuk synon të jetë platformë e plotë juridike për pajtueshmëri. Ai është prototip funksional që demonstron funksionet bazë dhe mënyrën e integrimit të tyre në një NVM. Për këtë arsye, operacionet si procesimi i ruajtjes së të dhënave dhe përfundimi i kërkesës për fshirje kontrollohen manualisht nga administratori.')

heading('3.2 Aktorët e sistemit', 2)
para('Aktorët kryesorë të sistemit janë Administratori dhe Përdoruesi. Administratori përfaqëson një person të autorizuar nga biznesi, për shembull pronarin, menaxherin ose personin përgjegjës për administrimin e klientëve dhe privatësisë. Përdoruesi përfaqëson një klient ose një person që ka llogari në aplikacion dhe dëshiron të shikojë ose të kontrollojë mënyrën se si trajtohen të dhënat e tij.')
para('Administratori mund të shikojë dashboard-in, të krijojë dhe të ndryshojë politika të privatësisë, të krijojë versione të reja, të shikojë miratimet, të filtrojë audit logs, të administron politikat e ruajtjes, të shqyrtojë kërkesat për të dhëna dhe të menaxhojë klientët e aplikacionit demonstrues. Ai gjithashtu mund të miratojë, të refuzojë ose të shënojë si të përfunduar kërkesat e përdoruesve.')
para('Përdoruesi ka qasje të kufizuar në funksionet që lidhen me privatësinë e tij. Ai mund të shikojë politikën aktuale, të pranojë ose të tërheqë miratimin dhe të kërkojë eksportim ose fshirje të të dhënave. Kjo ndarje e roleve paraqet një kërkesë të rëndësishme sigurie, sepse një përdorues i zakonshëm nuk duhet të mund të ndryshojë politikat e biznesit ose të shikojë të dhënat e klientëve të tjerë.')

heading('3.3 Kërkesat funksionale', 2)
para('Kërkesat funksionale përshkruajnë veprimet që sistemi duhet të jetë në gjendje t’i kryejë. Ato janë nxjerrë nga objektivi i punimit, nga rolet e identifikuara dhe nga rrjedhat që duhet të demonstrohen në aplikacion. Tabela 1 paraqet kërkesat kryesore funksionale të DataGuard.')

caption = doc.add_paragraph(); caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; caption.paragraph_format.space_after = Pt(4); caption.paragraph_format.first_line_indent = Cm(0); rr = caption.add_run('Tabela 1. Kërkesat funksionale të sistemit'); run_font(rr, bold=True, size=11)
rows = [
    ('KF-01', 'Autentifikimi', 'Sistemi duhet të lejojë krijimin e llogarisë, kyçjen dhe daljen e përdoruesve.'),
    ('KF-02', 'Rolet', 'Sistemi duhet të kufizojë funksionet sipas rolit Admin ose User.'),
    ('KF-03', 'Politikat', 'Administratori duhet të krijojë, shikojë dhe ndryshojë politikat e privatësisë.'),
    ('KF-04', 'Versionimi', 'Administratori duhet të krijojë version të ri të politikës duke ruajtur historikun.'),
    ('KF-05', 'Miratimet', 'Përdoruesi duhet të pranojë ose të tërheqë miratimin për politikën aktuale.'),
    ('KF-06', 'Auditimi', 'Sistemi duhet të ruajë dhe të shfaqë veprimet e rëndësishme me datë, aktor dhe status.'),
    ('KF-07', 'Ruajtja', 'Administratori duhet të krijojë dhe të ndryshojë politikat e ruajtjes së të dhënave.'),
    ('KF-08', 'Kërkesat e të dhënave', 'Përdoruesi duhet të kërkojë eksportim ose fshirje, ndërsa administratori duhet t’i procesojë kërkesat.'),
    ('KF-09', 'Klientët', 'Administratori duhet të shtojë, shikojë, ndryshojë dhe fshijë klientë në aplikacionin SME.'),
    ('KF-10', 'Dashboard', 'Administratori duhet të shohë statistika të përmbledhura dhe aktivitetet e fundit.')
]
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
for c, t in zip(table.rows[0].cells, ['ID', 'Funksioni', 'Përshkrimi']): set_cell_text(c, t, True)
for row in rows:
    cells = table.add_row().cells
    for cell, text in zip(cells, row): set_cell_text(cell, text)
set_table_geometry(table, [900, 2100, 6360])

para('Kërkesat funksionale janë implementuar si faqe dhe komponentë të veçantë në ndërfaqe. Për shembull, faqja Privacy Policies mbulon KF-03 dhe KF-04, faqja My Privacy mbulon KF-05 dhe një pjesë të KF-08, ndërsa Audit Logs mbulon KF-06. Kjo ndarje e bën sistemin të kuptueshëm dhe mundëson testimin e secilës kërkesë në mënyrë të pavarur.')

heading('3.4 Kërkesat jofunksionale', 2)
para('Kërkesat jofunksionale përshkruajnë cilësitë që duhet të ketë sistemi, jo vetëm funksionet që kryen. Kërkesa e parë është përdorshmëria. Ndërfaqja duhet të jetë e pastër, e kuptueshme dhe e përshtatshme për përdorues që nuk kanë njohuri të avancuara teknike. Për këtë arsye, menuja anësore ndan funksionet sipas modulit, ndërsa butonat dhe formularët përdorin emërtime të qarta.')
para('Kërkesa e dytë është siguria bazë e qasjes. Sistemi përdor Supabase Auth për autentifikim dhe ruan rolin e përdoruesit në tabelën profiles. Rregullat Row Level Security përdoren për të kontrolluar që përdoruesi të mund të qaset vetëm te të dhënat e veta, ndërsa administratori të ketë qasje në funksionet e menaxhimit. Në nivel të demonstrimit, kjo ndarje është e mjaftueshme për të treguar parimin e autorizimit me role.')
para('Kërkesa e tretë është ruajtja e qëndrueshme e të dhënave. Informacioni për politikat, miratimet, kërkesat dhe klientët duhet të mbetet i ruajtur pas rifreskimit të aplikacionit. Për këtë arsye përdoret PostgreSQL përmes Supabase. Kërkesa e katërt është gjurmueshmëria, që nënkupton ruajtjen e veprimeve të rëndësishme në audit logs dhe mundësinë e filtrimit të tyre sipas datës dhe veprimit.')
para('Kërkesat e tjera jofunksionale janë përgjegjshmëria e ndërfaqes, mirëmbajtja e kodit dhe performanca e pranueshme për një numër të vogël ose mesatar të regjistrimeve. Aplikacioni përdor React, TypeScript dhe komponentë të ripërdorshëm për ta bërë kodin më të organizuar. Nuk është bërë testim i ngarkesës në shkallë të madhe, sepse kjo tejkalon fushëveprimin e një prototipi Bachelor.')

heading('3.5 Rastet e përdorimit', 2)
para('Rasti i parë i përdorimit është krijimi i politikës së privatësisë. Administratori kyçet në sistem, hap faqen Privacy Policies dhe zgjedh “Create Policy”. Ai jep titullin, versionin, datën e hyrjes në fuqi dhe tekstin e politikës. Pas ruajtjes, politika bëhet e dukshme në listën administrative dhe mund të shikohet nga përdoruesi në seksionin My Privacy.')
para('Rasti i dytë është pranimi i miratimit. Përdoruesi kyçet në llogarinë e tij, hap My Privacy, lexon politikën aktuale dhe zgjedh “Accept Privacy Policy”. Sistemi krijon një regjistrim të miratimit që përfshin identitetin e përdoruesit, versionin e politikës, statusin Accepted dhe datën e vendimit. Nëse përdoruesi e ndryshon vendimin, zgjedh “Withdraw Consent”, ndërsa sistemi ruan statusin Withdrawn.')
para('Rasti i tretë është kërkesa për të dhëna. Përdoruesi zgjedh “Request & Download My Data” ose “Request Data Deletion”. Për eksportim, sistemi krijon kërkesë dhe mundëson shkarkimin e një skedari JSON me të dhënat e disponueshme në prototip. Për fshirje, krijohet kërkesë me status Pending. Administratori e hap faqen Data Requests, e aprovon ose e refuzon kërkesën dhe, kur procesi mbaron, e shënon Completed.')
para('Rasti i katërt lidhet me aplikacionin demonstrues SME. Administratori hap SME Customer App, shton një klient të ri ose ndryshon të dhënat e një klienti ekzistues. Veprimet e rëndësishme regjistrohen në audit logs. Kështu demonstrohet lidhja ndërmjet modulit të privatësisë dhe një aplikacioni që ruan të dhëna personale të klientëve.')

heading('3.6 Kufizimet e sistemit', 2)
para('DataGuard është ndërtuar si prototip i kontrolluar dhe ka disa kufizime të qëllimshme. Së pari, sistemi mbështet vetëm dy role dhe nuk përfshin nivele të shumta autorizimi, menaxhim të ekipeve ose organizatave të shumta. Së dyti, politikat e ruajtjes paraqiten dhe procesohen manualisht; nuk ekziston mekanizëm automatik që i fshin ose i arkivon të dhënat kur mbaron periudha e ruajtjes.')
para('Së treti, eksporti i të dhënave është i thjeshtuar dhe përmban informacionin e disponueshëm për përdoruesin në format JSON. Nuk realizohet verifikim shtesë i identitetit, enkriptim i skedarit të eksportuar ose dërgim me email. Së katërti, kërkesa për fshirje menaxhohet si rrjedhë administrative dhe nuk përfshin anonimizim automatik, kontroll të detajuar të afateve ligjore ose lidhje me burime të tjera të të dhënave.')
para('Këto kufizime nuk paraqesin dështim të sistemit, por kufizim të fushëveprimit. Ato janë zgjedhur për të ruajtur thjeshtësinë e prototipit dhe për t’u përqendruar në objektivat kryesore të temës: miratimet, versionimi, auditimi, ruajtja dhe kërkesat për të dhëna. Në punë të ardhshme, sistemi mund të zgjerohet me njoftime, automatizim të ruajtjes, role shtesë, anonimizim dhe integrime me aplikacione të tjera SME.')

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run('DataGuard - Punim diplome'); run_font(r, size=10)
OUT.parent.mkdir(exist_ok=True); doc.save(OUT); print(OUT)
