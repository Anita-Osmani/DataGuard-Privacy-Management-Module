from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUT = Path('outputs/Kapitulli_2_Konteksti_Ligjor_dhe_Literatura.docx')
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(2.5), Cm(2.5)
section.left_margin, section.right_margin = Cm(3), Cm(3)
section.header_distance, section.footer_distance = Cm(1.25), Cm(1.25)

normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
normal._element.rPr.rFonts.set(qn('w:cs'), 'Times New Roman')
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.first_line_indent = Cm(1.25)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for name, size in [('Heading 1', 14), ('Heading 2', 12)]:
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
    style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    style.font.size, style.font.bold = Pt(size), True
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(12), Pt(6)
    style.paragraph_format.first_line_indent = Cm(0)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

def heading(text, level): doc.add_paragraph(text, style=f'Heading {level}')
def para(text):
    p = doc.add_paragraph(style='Normal')
    p.add_run(text)
    return p

heading('2. KONTEKSTI LIGJOR DHE RISHIKIMI I LITERATURËS', 1)

heading('2.1 Privatësia dhe mbrojtja e të dhënave', 2)
para('Privatësia paraqet të drejtën e individit për të pasur kontroll të arsyeshëm mbi informacionin që lidhet me jetën e tij personale. Në mjedisin digjital, kjo e drejtë lidhet ngushtë me mbledhjen, ruajtjen, përdorimin, ndarjen dhe fshirjen e të dhënave personale. Mbrojtja e të dhënave personale është shndërruar në një pjesë të rëndësishme të zhvillimit të sistemeve informative, sepse pothuajse çdo shërbim digjital ruan të paktën një formë të të dhënës që mund të lidhet me një person fizik.')
para('Dallimi ndërmjet privatësisë dhe mbrojtjes së të dhënave është i dobishëm për kuptimin e këtij punimi. Privatësia është koncept më i gjerë dhe lidhet me respektimin e jetës personale e familjare, ndërsa mbrojtja e të dhënave përqendrohet në rregullat dhe masat që zbatohen gjatë përpunimit të informacionit personal. Në praktikë, një aplikacion që respekton privatësinë duhet t’i japë përdoruesit informacion të qartë për përpunimin e të dhënave, të kufizojë mbledhjen e panevojshme të tyre dhe të mundësojë ushtrimin e të drejtave bazë.')
para('Ligji i Kosovës nr. 06/L-082 për Mbrojtjen e të Dhënave Personale përcakton të drejtat, përgjegjësitë, parimet dhe masat ndëshkuese që lidhen me mbrojtjen e të dhënave personale. Ligji e lidh mbrojtjen e të dhënave me privatësinë e individit dhe është i harmonizuar me Rregulloren (BE) 2016/679 [1]. Për mbikëqyrjen e zbatimit të këtij ligji përgjegjëse është Agjencia për Informim dhe Privatësi, e cila vepron si autoritet i pavarur mbikëqyrës [1].')
para('Për NVM-të, mbrojtja e të dhënave nuk duhet të kuptohet vetëm si çështje ligjore. Ajo ndikon në besimin e klientëve, cilësinë e proceseve të brendshme dhe aftësinë e biznesit për të administruar informacionin në mënyrë të rregullt. Një zgjidhje e thjeshtë teknike, e integruar në aplikacionin ekzistues, mund të ndihmojë biznesin të kalojë nga ruajtja e paorganizuar e të dhënave drejt proceseve më të qarta dhe më të dokumentuara.')

heading('2.2 Të dhënat personale dhe përpunimi i tyre', 2)
para('Të dhëna personale konsiderohen të gjitha informacionet që lidhen me një person fizik të identifikuar ose të identifikueshëm. Emri, mbiemri, adresa elektronike, numri i telefonit, adresa e banimit dhe identifikuesit online janë shembuj të zakonshëm. Në një aplikacion për klientë të NVM-ve, kombinimi i emrit me kontaktet ose adresën e klientit krijon qartë një grup të dhënash personale që kërkon trajtim të kujdesshëm.')
para('Përpunimi i të dhënave personale përfshin një gamë të gjerë veprimesh: mbledhjen, regjistrimin, organizimin, ruajtjen, ndryshimin, shikimin, përdorimin, dërgimin, kufizimin, fshirjen ose shkatërrimin e të dhënave. Ligji i Kosovës e përkufizon përpunimin në këtë mënyrë të gjerë [1]. Kjo do të thotë se edhe një veprim i thjeshtë, si shikimi i profilit të klientit nga administratori, mund të jetë relevant nga aspekti i kontrollit të qasjes dhe auditimit.')
para('Në këtë punim, aplikacioni demonstrues ruan katër kategori bazë të të dhënave të klientëve: emrin, emailin, telefonin dhe adresën. Zgjedhja është e qëllimshme, sepse këto të dhëna janë të zakonshme për një biznes të vogël dhe e bëjnë të qartë lidhjen ndërmjet aplikacionit të klientëve dhe modulit të privatësisë. Prototipi nuk përpunon kategori të veçanta të të dhënave, si të dhënat shëndetësore, biometrike ose të dhënat për bindjet politike; këto kategori do të kërkonin masa dhe analiza shtesë.')

heading('2.3 Parimet e mbrojtjes së të dhënave', 2)
para('Parimet e përpunimit të të dhënave krijojnë bazën për projektimin e një sistemi privatësie. GDPR përcakton parimet e ligjshmërisë, drejtësisë dhe transparencës; kufizimit të qëllimit; minimizimit të të dhënave; saktësisë; kufizimit të ruajtjes; integritetit dhe konfidencialitetit; si dhe llogaridhënies [2]. Këto parime nuk janë vetëm formulime juridike, por mund të përkthehen në vendime konkrete gjatë projektimit të një aplikacioni.')
para('Parimi i transparencës lidhet me paraqitjen e një politike të qartë të privatësisë. Në DataGuard, përdoruesi mund të shikojë politikën aktuale dhe versionin e saj. Kufizimi i qëllimit dhe minimizimi i të dhënave reflektohen në ruajtjen e vetëm të fushave që kërkohen për demonstrimin e menaxhimit të klientëve. Kufizimi i ruajtjes paraqitet përmes politikave të ruajtjes, ku për secilën kategori përcaktohet një periudhë dhe një veprim, për shembull fshirje, arkivim ose ruajtje.')
para('Llogaridhënia kërkon që organizata të jetë në gjendje të tregojë se ka zbatuar masa të arsyeshme. Në një prototip të thjeshtë, audit logs dhe historiku i versioneve ofrojnë dëshmi praktike për këtë qëllim. Ato nuk e zëvendësojnë një proces të plotë të pajtueshmërisë ligjore, por krijojnë një bazë të dobishme për gjurmueshmëri dhe kontroll të brendshëm.')

heading('2.4 GDPR dhe rëndësia për NVM-të', 2)
para('GDPR është Rregullorja e Përgjithshme e Bashkimit Evropian për Mbrojtjen e të Dhënave. Ajo vendos kërkesa për përpunimin e të dhënave personale dhe ka ndikuar në praktikat e organizatave që ofrojnë shërbime digjitale. Rregullorja thekson se përpunimi duhet të ketë bazë të ligjshme dhe se subjekti i të dhënave duhet të informohet në mënyrë të qartë [2]. Ligji i Kosovës nr. 06/L-082 është i harmonizuar me këtë rregullore, çka e bën GDPR një referencë të rëndësishme edhe në kontekstin vendor [1].')
para('Për NVM-të, sfida nuk është domosdoshmërisht mungesa e vullnetit për të respektuar privatësinë, por përkthimi i kërkesave të përgjithshme në funksione të përdorshme. Një biznes i vogël nuk ka gjithmonë ekip juridik, specialist për privatësi ose buxhet për platforma të mëdha. Për këtë arsye, një modul bazë duhet të fokusohet në funksionet me vlerë të drejtpërdrejtë: politikë e qartë, miratim i regjistruar, histori versionesh, auditim, periudha ruajtjeje dhe rrjedha për kërkesat e përdoruesve.')
para('GDPR parashikon të drejtën e qasjes në të dhënat personale, të drejtën për fshirje në rrethana të caktuara dhe të drejtën për pranimin e një kopjeje të të dhënave në format të zakonshëm elektronik [2]. Në DataGuard këto të drejta demonstrohen në nivel prototipi përmes faqes “My Privacy”, ku përdoruesi mund të kërkojë dhe të shkarkojë të dhënat e veta ose të paraqesë kërkesë për fshirje. Vendimi përfundimtar dhe procesimi i kërkesës kryhen nga administratori.')

heading('2.5 Menaxhimi i miratimeve të përdoruesve', 2)
para('Miratimi është një nga bazat e mundshme ligjore për përpunimin e të dhënave personale. Ai nuk duhet të trajtohet si një kuti e parazgjedhur ose si një veprim i paqartë. Udhëzimet e Bordit Evropian për Mbrojtjen e të Dhënave theksojnë rëndësinë e miratimit të vlefshëm dhe të dëshmueshëm, si dhe detyrën e kontrolluesit për të vlerësuar bazën e përshtatshme ligjore për përpunimin e parashikuar [3].')
para('Nga perspektiva e sistemit, menaxhimi i miratimeve kërkon lidhjen e qartë ndërmjet përdoruesit, politikës së privatësisë, versionit të politikës, statusit të vendimit dhe datës së vendimit. Nëse politika ndryshon, duhet të ekzistojë mundësia që të krijohet version i ri, ndërsa versionet e mëparshme të mos humbin. Kjo lejon që miratimi të mos lidhet vetëm me titullin e një politike, por me tekstin konkret që ka qenë i disponueshëm në kohën e pranimit.')
para('DataGuard implementon statuset Accepted, Rejected dhe Withdrawn për regjistrimet e miratimeve. Në ndërfaqen e përdoruesit, përdoruesi mund të pranojë politikën aktuale ose të tërheqë miratimin. Administratori mund të shikojë regjistrat e miratimeve së bashku me politikën, versionin dhe datën. Ky model është i kufizuar qëllimisht në një politikë kryesore privatësie dhe nuk synon menaxhimin e avancuar të preferencave për marketing, cookies, gjuhë të shumta ose pajtueshmëri për juridiksione të ndryshme.')

heading('2.6 Ruajtja, eksportimi dhe fshirja e të dhënave', 2)
para('Ruajtja e pakufizuar e të dhënave rrit rrezikun dhe e bën më të vështirë menaxhimin e informacionit. Parimi i kufizimit të ruajtjes kërkon që të dhënat të mbahen vetëm për periudhën e nevojshme për qëllimin e përpunimit [2]. Në mjediset reale, periudha përcaktohet duke marrë parasysh natyrën e të dhënave, qëllimin e biznesit, nevojat operative dhe kërkesat ligjore.')
para('Politikat e ruajtjes shërbejnë për ta bërë këtë vendim të dukshëm dhe të dokumentuar. Në DataGuard, administratorit i paraqiten politika për kategori të tilla si Customer Data, Audit Logs dhe Consent Records. Për secilën kategori mund të ruhet periudha e ruajtjes, veprimi i parashikuar dhe statusi. Butoni “Process Retention” përfaqëson një proces manual të shqyrtimit; sistemi nuk kryen fshirje automatike në sfond. Ky kufizim është i përshtatshëm për një prototip Bachelor, sepse e mban zgjidhjen të kuptueshme dhe shmang fshirjet e pakontrolluara.')
para('Eksportimi i të dhënave lidhet me mundësinë e përdoruesit për të marrë kopjen e informacionit të tij. Në prototip, përdoruesi mund të shkarkojë një dokument JSON që përmbledh profilin, miratimet dhe kërkesat e veta. Kërkesa për fshirje regjistrohet në sistem dhe kalon nëpër statuset Pending, Approved, Completed ose Rejected. Këto statuse e ndajnë paraqitjen e kërkesës nga vendimi administrativ dhe nga përfundimi i procesit.')

heading('2.7 Regjistrimi i qasjeve dhe auditimi', 2)
para('Auditimi është procesi i regjistrimit dhe shqyrtimit të veprimeve të rëndësishme të kryera në një sistem. Në kontekstin e privatësisë, audit logs ndihmojnë që organizata të kuptojë se çfarë është bërë me të dhënat, kur është bërë, nga kush dhe me çfarë rezultati. Një regjistër i tillë mbështet kontrollin e brendshëm, analizën e gabimeve dhe përgjegjësinë e përdoruesve me qasje administrative.')
para('Struktura bazë e një audit log përfshin datën, përdoruesin ose aktorin, veprimin, burimin dhe statusin. Në DataGuard përdoren veprime si LOGIN, VIEW, CREATE, UPDATE, DELETE, CONSENT dhe EXPORT. Për shembull, krijimi i një politike, ndryshimi i klientit ose procesimi manual i politikave të ruajtjes mund të regjistrohen si aktivitete. Administratori mund t’i filtrojë regjistrimet sipas veprimit dhe datës.')
para('Auditimi nuk duhet të ngatërrohet me një mekanizëm të pandryshueshëm blockchain ose me monitorim të komplikuar të sigurisë. Për qëllimin e këtij punimi, mjafton ruajtja e një historie të thjeshtë dhe të lexueshme të veprimeve të rëndësishme. Ky dizajn përputhet me objektivin e projektit: të ofrojë një bazë praktike për NVM-të, jo një platformë të nivelit enterprise.')

heading('2.8 Rishikimi i zgjidhjeve ekzistuese', 2)
para('Ekzistojnë platforma komerciale që ofrojnë menaxhim të avancuar të privatësisë, miratimeve dhe kërkesave të subjekteve të të dhënave. Një shembull është OneTrust Consent & Preferences, e cila paraqet mjete për transparencë, kontroll të preferencave, ruajtje të regjistrimeve të miratimeve, histori ndryshimesh dhe logje të eksportueshme [5]. Zgjidhje të këtij lloji zakonisht mbulojnë shumë kanale, rajone, juridiksione dhe integrime me sisteme të tjera.')
para('Megjithatë, platformat e tilla janë të dizajnuara shpesh për organizata me kërkesa të gjera, volume të larta të të dhënave dhe procese të ndërlikuara të pajtueshmërisë. Ato mund të përfshijnë zbulim automatik të cookies, kategorizim të aseteve, menaxhim shumëgjuhësh, raportim të avancuar dhe automatizim të kërkesave. Këto veçori janë të dobishme, por nuk janë të domosdoshme për një prototip të orientuar për NVM-të.')
para('NIST Privacy Framework ofron një qasje vullnetare për identifikimin dhe menaxhimin e rreziqeve të privatësisë. Korniza synon të ndihmojë organizatat të menaxhojnë rrezikun e privatësisë gjatë krijimit të produkteve dhe shërbimeve [4]. Në këtë punim, ideja kryesore e kësaj qasjeje është përkthyer në funksione konkrete: kontrolli i miratimit, komunikimi përmes politikës, mbrojtja e qasjes me role dhe regjistrimi i aktiviteteve.')
para('DataGuard zë një pozitë ndërmjet proceseve tërësisht manuale dhe platformave të mëdha enterprise. Ai nuk synon të konkurrojë me zgjidhjet komerciale, por të demonstrojë një arkitekturë të thjeshtë, të kuptueshme dhe të zgjerueshme. Përparësia e zgjidhjes është fokusimi në funksionet bazë që mund të integrohen në një aplikacion SME: politikat dhe versionet, miratimet, auditimi, ruajtja, kërkesat për të dhëna dhe një aplikacion demonstrues për klientët.')

heading('Referenca të përdorura në këtë kapitull', 2)
refs = [
    '[1] Kuvendi i Republikës së Kosovës. Ligji nr. 06/L-082 për Mbrojtjen e të Dhënave Personale, Gazeta Zyrtare nr. 6/2019, 25 shkurt 2019. Në dispozicion: https://gzk.rks-gov.net/ActDocumentDetail.aspx?ActID=18616',
    '[2] European Parliament and Council of the European Union. Regulation (EU) 2016/679 (General Data Protection Regulation), 27 April 2016. EUR-Lex. Në dispozicion: https://eur-lex.europa.eu/eli/reg/2016/679/oj',
    '[3] European Data Protection Board. Guidelines 05/2020 on consent under Regulation 2016/679, Version 1.1, 4 May 2020. Në dispozicion: https://www.edpb.europa.eu/our-work-tools/our-documents/guidelines/guidelines-052020-consent-under-regulation-2016679_en',
    '[4] National Institute of Standards and Technology. NIST Privacy Framework, Version 1.0, January 2020. Në dispozicion: https://www.nist.gov/privacy-framework/privacy-framework',
    '[5] OneTrust. Consent Management Platform. Në dispozicion: https://www.onetrust.com/products/consent-management/'
]
for ref in refs:
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(3)
    p.add_run(ref)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run('DataGuard - Punim diplome')
r.font.name = 'Times New Roman'; r._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman'); r.font.size = Pt(10)
OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(OUT)
