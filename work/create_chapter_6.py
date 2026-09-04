from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUT = Path('outputs/Kapitulli_6_Regjistrimi_i_Qasjeve_dhe_Auditimi.docx')


def set_font(run, size=12, bold=False):
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    run._element.rPr.rFonts.set(qn('w:cs'), 'Times New Roman')
    run.font.size = Pt(size)
    run.font.bold = bold


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.5)
sec.left_margin, sec.right_margin = Cm(3), Cm(3)
sec.header_distance, sec.footer_distance = Cm(1.25), Cm(1.25)

normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
normal._element.rPr.rFonts.set(qn('w:cs'), 'Times New Roman')
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.first_line_indent = Cm(1.25)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for style_name, size in [('Heading 1', 14), ('Heading 2', 12)]:
    style = doc.styles[style_name]
    style.font.name = 'Times New Roman'
    style._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
    style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    style.font.size = Pt(size)
    style.font.bold = True
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.first_line_indent = Cm(0)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


def heading(text, level):
    doc.add_paragraph(text, style=f'Heading {level}')


def para(text):
    p = doc.add_paragraph(style='Normal')
    p.add_run(text)
    return p


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def no_indent(cell):
    for p in cell.paragraphs:
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.1
        p.paragraph_format.space_after = Pt(0)


heading('6. MODULI I REGJISTRIMIT TË QASJEVE DHE AUDITIMIT', 1)

heading('6.1 Qëllimi i auditimit', 2)
para('Regjistrimi i qasjeve dhe auditimi janë pjesë e rëndësishme e një moduli të privatësisë, sepse mundësojnë gjurmimin e veprimeve që kryhen mbi të dhënat dhe funksionet e aplikacionit. Në kontekstin e një NVM-je, auditimi nuk kërkon domosdoshmërisht një infrastrukturë të ndërlikuar. Një regjistër i thjeshtë, i ruajtur në bazën e të dhënave dhe i paraqitur në panelin administrativ, ndihmon biznesin të kuptojë kush ka kryer një veprim, çfarë veprimi është kryer, mbi cilin burim dhe kur ka ndodhur ai veprim.')
para('Qëllimi kryesor i modulit Audit Logs në DataGuard është përgjegjshmëria. Kur administratori krijon një politikë, përditëson të dhënat e një klienti, përpunon kërkesë për fshirje ose eksporton të dhëna, veprimi përkatës mund të shfaqet në historikun e auditimit. Ky historik është i dobishëm për kontroll të brendshëm, për sqarimin e gabimeve operative dhe për demonstrimin e rrjedhës së administrimit të të dhënave personale.')
para('Sipas udhëzimeve të OWASP, regjistrat e aplikacionit janë të dobishëm si për qëllime operative ashtu edhe për siguri, ndërsa audit trail mund të përfshijë shtimin, ndryshimin, fshirjen dhe eksportimin e të dhënave [1]. Kjo ide është përdorur si parim orientues në DataGuard: regjistrohen veprime të rëndësishme të biznesit dhe të privatësisë, por nuk ruhet çdo klikim i përdoruesit. Kështu shmanget krijimi i një sasie të panevojshme informacioni në regjistra.')
para('Prototipi nuk implementon blockchain, hash-chain, analizë automatike të anomalive ose monitorim në kohë reale. Këto funksione nuk janë të domosdoshme për objektivin e punimit dhe do ta bënin zgjidhjen më të ndërlikuar se nevojat e një demonstrimi për NVM. Në vend të tyre, është zgjedhur një model i kuptueshëm me tabelën audit_logs, kontroll të qasjes dhe filtra bazë për leximin e regjistrimeve.')

heading('6.2 Regjistrimi i aktiviteteve', 2)
para('Në DataGuard, aktivitetet e auditueshme ndahen sipas llojit të veprimit. Vlerat kryesore të përdorura janë LOGIN, VIEW, CREATE, UPDATE, DELETE, CONSENT dhe EXPORT. LOGIN paraqet hyrjen në sistem; VIEW paraqet shikimin e një burimi; CREATE, UPDATE dhe DELETE lidhen me krijimin, përditësimin ose largimin e të dhënave; CONSENT lidhet me pranimin ose tërheqjen e miratimit; ndërsa EXPORT lidhet me përgatitjen e një eksporti të të dhënave personale.')
para('Në implementimin e prototipit, auditimi lidhet veçanërisht me veprimet administrative me ndikim mbi të dhënat. Shembuj të tillë janë krijimi ose përditësimi i politikës së privatësisë, krijimi i versionit të ri, përpunimi i kërkesës për të dhëna, krijimi, ndryshimi dhe fshirja e klientit në aplikacionin demonstrues SME, si dhe përpunimi manual i politikave të ruajtjes. Pas përfundimit të veprimit, aplikacioni krijon një regjistrim auditimi dhe më pas rifreskon listën që paraqitet në panel.')
para('Për përdoruesin e zakonshëm, rrjedhat kryesore janë shikimi i politikës, pranimi ose tërheqja e miratimit dhe krijimi i kërkesës për eksport ose fshirje. Këto rrjedha janë ndarë nga administrimi i audit logs, sepse përdoruesi nuk ka nevojë të lexojë historikun e plotë të veprimeve të biznesit. Ky ndalim i qasjes është në përputhje me parimin e minimizimit të informacionit: çdo rol sheh vetëm funksionet që i nevojiten.')
para('Zgjedhja e aktiviteteve për regjistrim duhet të jetë proporcionale me rrezikun dhe me qëllimin e sistemit. OWASP thekson se aplikacioni ka informacion për identitetin e përdoruesit dhe kontekstin e veprimit, si veprimin, objektin dhe rezultatin [1]. Në DataGuard, ky informacion reduktohet në fushat e nevojshme për një prototip: përdoruesi që e ka kryer veprimin, tipi i veprimit, burimi i prekur, statusi dhe koha e regjistrimit.')

heading('6.3 Struktura e audit logs', 2)
para('Regjistrat ruhen në tabelën audit_logs në PostgreSQL përmes Supabase. Çdo rresht përfaqëson një ngjarje të vetme auditimi. Tabela lidhet me profiles përmes fushës actor_id, e cila tregon përdoruesin që ka kryer veprimin. Kjo lidhje e bën të mundur që në panel të paraqitet emri i përdoruesit në vend të një identifikuesi teknik.')
para('Struktura e tabelës është mbajtur e thjeshtë, sepse qëllimi nuk është ruajtja e të dhënave të plota të klientit në regjistër. Për shembull, kur krijohet një klient, në audit log ruhet fakti se është kryer CREATE mbi burimin Customer; nuk ruhet e gjithë adresa ose numri i telefonit i klientit. Kjo zvogëlon kopjimin e të dhënave personale në një tabelë shtesë dhe e bën historikun më të përqendruar në veprim.')

table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = 'Table Grid'
# Keep the Word table geometry explicit and equal to the printable width.
tbl_pr = table._tbl.tblPr
tbl_w = tbl_pr.first_child_found_in('w:tblW')
tbl_w.set(qn('w:type'), 'dxa')
tbl_w.set(qn('w:w'), '8505')
tbl_ind = OxmlElement('w:tblInd')
tbl_ind.set(qn('w:w'), '120')
tbl_ind.set(qn('w:type'), 'dxa')
tbl_pr.append(tbl_ind)
headers = ['Fusha', 'Përshkrimi', 'Qëllimi në auditim']
for index, text in enumerate(headers):
    cell = table.rows[0].cells[index]
    cell.text = text
    shade(cell, 'E6E6E6')
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    no_indent(cell)
    for run in cell.paragraphs[0].runs:
        set_font(run, 11, True)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

rows = [
    ('id', 'Identifikues unik i regjistrimit.', 'Dallon çdo ngjarje auditimi.'),
    ('actor_id', 'Referencë te profili i përdoruesit që kreu veprimin.', 'Tregon se kush e kreu veprimin.'),
    ('action', 'Lloji i veprimit, p.sh. CREATE, UPDATE ose EXPORT.', 'Tregon çfarë ndodhi.'),
    ('resource', 'Burimi i prekur, p.sh. Customer, Policy ose Data Request.', 'Tregon objektin e veprimit.'),
    ('status', 'Rezultati i veprimit, zakonisht Success.', 'Tregon nëse veprimi u realizua.'),
    ('created_at', 'Data dhe koha e krijimit të regjistrit.', 'Tregon kur ndodhi veprimi.')
]
for field, description, purpose in rows:
    cells = table.add_row().cells
    for i, text in enumerate((field, description, purpose)):
        cells[i].text = text
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        no_indent(cells[i])
        for run in cells[i].paragraphs[0].runs:
            set_font(run, 10.5, i == 0)

caption = doc.add_paragraph()
caption.paragraph_format.first_line_indent = Cm(0)
caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = caption.add_run('Tabela 6.1. Fushat kryesore të tabelës audit_logs')
set_font(r, 11, False)
r.italic = True

para('Fushat action, resource dhe status e formojnë përshkrimin bazë të ngjarjes. Këto fusha përputhen me pyetjet praktike që administratori dëshiron të bëjë gjatë kontrollit: çfarë ndodhi, mbi cilin burim dhe me çfarë rezultati. Data dhe koha ruhen përmes created_at, ndërsa lidhja me profilin e përdoruesit ruhet përmes actor_id. Udhëzimi i OWASP përmbledh një audit record të dobishëm si informacion për "kur, ku, kush dhe çfarë"; DataGuard mbulon pjesët më të rëndësishme të këtij parimi për funksionin e tij [1].')
para('Për sigurinë e të dhënave, nuk rekomandohet të ruhen fjalëkalime, tokenë, informacione të plota të kartelave ose tekste të panevojshme të formularëve në audit log. Regjistri duhet të japë kontekst të mjaftueshëm për verifikim, por jo të krijojë një kopje të dytë të të dhënave të ndjeshme. Në një version të ardhshëm, mund të shtohen fusha të kontrolluara si arsyeja e dështimit ose identifikuesi i kërkesës, pa ruajtur përmbajtje të panevojshme personale.')

heading('6.4 Shfaqja dhe filtrimi i regjistrimeve', 2)
para('Faqja Audit Logs është e disponueshme vetëm për administratorin. Ajo paraqet regjistrimet në formë tabele me kolonat Date, User, Action, Resource dhe Status. Renditja e regjistrimeve nga më e reja tek më e vjetra i jep administratorit pamje të shpejtë mbi aktivitetin e fundit të sistemit. Në Dashboard paraqitet gjithashtu një përmbledhje e aktiviteteve të fundit, e cila shërben si hyrje e shpejtë në auditim.')
para('Filtrat janë qëllimisht të thjeshtë. Administratori mund të zgjedhë një veprim të caktuar, si CREATE ose DELETE, ose të shfaqë të gjitha veprimet. Po ashtu, mund të kufizojë rezultatet sipas datës. Kombinimi i filtrit të veprimit me filtrin e datës ndihmon në një skenar praktik, për shembull kur kërkohet të shihet kush ka përpunuar kërkesat për të dhëna në një ditë të caktuar ose kur duhet të kontrollohen vetëm fshirjet e klientëve.')
para('Filtrimi kryhet në ndërfaqe mbi të dhënat e lexuara nga tabela audit_logs. Për një prototip me numër të vogël regjistrimesh, kjo zgjidhje është e mjaftueshme dhe e kuptueshme. Nëse aplikacioni do të përdorej nga shumë biznese ose do të krijonte mijëra regjistrime, filtrimi, faqëzimi dhe kërkimi do të duhej të optimizoheshin në nivel të bazës së të dhënave ose përmes një API-je të dedikuar.')
para('Qasja në audit logs kontrollohet me role dhe me Row Level Security të Supabase. Roli Admin ka qasje në menaxhimin e regjistrimeve, ndërsa roli User nuk merr qasje në faqen e auditimit. Ky ndalim është i rëndësishëm sepse vetë regjistrat mund të përmbajnë informacion mbi aktivitetet e përdoruesve të tjerë. Prandaj, auditimi nuk është vetëm proces i ruajtjes së log-eve, por edhe proces i kontrollit se kush mund t’i lexojë ato.')

heading('6.5 Implementimi dhe rezultatet', 2)
para('Në anën e aplikacionit, DataGuard përdor një funksion ndihmës për regjistrimin e veprimeve të rëndësishme. Pas një veprimi administrativ, funksioni krijon një rresht të ri në audit_logs me actor_id të përdoruesit aktiv, action, resource dhe status. Pastaj lista e audit logs rifreskohet që administratori të mund ta shohë ngjarjen e re pa pasur nevojë të largohet nga faqja. Kjo e bën rrjedhën e auditimit të dukshme gjatë demonstrimit të sistemit.')
para('Testimi manual u krye duke u kyçur si administrator dhe duke kryer veprime në modulet kryesore. U verifikua se krijimi dhe përditësimi i politikave, veprimet mbi klientët dhe përpunimi i kërkesave mund të pasqyrohen në listën e auditimit me përdoruesin, veprimin, burimin, statusin dhe kohën. U testuan edhe filtrat e veprimit dhe të datës për të kontrolluar që lista mund të ngushtohet sipas nevojës së administratorit.')
para('Rezultati është një modul auditimi funksional për demonstrim, i cili ofron gjurmueshmëri bazë të aktiviteteve të rëndësishme. Ai mbështet objektivin e punimit duke e lidhur menaxhimin e privatësisë me një evidencë administrative të veprimeve. Rrjedha është e mjaftueshme për një NVM që dëshiron të kuptojë përdorimin e të dhënave brenda aplikacionit të saj, pa kërkuar një platformë të madhe monitorimi.')
para('Kufizimet e prototipit janë të qarta. Regjistrat nuk janë të pandryshueshëm në kuptimin kriptografik, nuk dërgohen në sistem të jashtëm monitorimi dhe nuk analizohet automatikisht sjellja e dyshimtë. Në një zgjerim të ardhshëm, sistemi mund të përfshijë ruajtje të centralizuar të log-eve, njoftime për veprime me rrezik të lartë, eksport të raportit të auditimit dhe politika të veçanta ruajtjeje për regjistrat. Këto zgjerime nuk janë pjesë e implementimit aktual, por tregojnë drejtimet e mundshme të zhvillimit.')

heading('Referenca të përdorura në këtë kapitull', 2)
references = [
    '[1] OWASP Foundation. Logging Cheat Sheet. Në dispozicion: https://cheatsheetseries.owasp.org/cheatsheets/Logging_Cheat_Sheet.html',
    '[2] European Parliament and Council of the European Union. Regulation (EU) 2016/679 (General Data Protection Regulation), 27 April 2016. EUR-Lex. Në dispozicion: https://eur-lex.europa.eu/eli/reg/2016/679/oj'
]
for reference in references:
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(3)
    p.add_run(reference)

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run('DataGuard - Punim diplome')
set_font(r, 10)

OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(OUT)
