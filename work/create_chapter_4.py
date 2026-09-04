from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path('outputs/Kapitulli_4_Arkitektura_dhe_Dizajni.docx')
ASSETS = Path('work/chapter4_assets')
ASSETS.mkdir(parents=True, exist_ok=True)

def font(size, bold=False):
    candidates = [r'C:\Windows\Fonts\arialbd.ttf' if bold else r'C:\Windows\Fonts\arial.ttf', r'C:\Windows\Fonts\calibrib.ttf' if bold else r'C:\Windows\Fonts\calibri.ttf']
    for path in candidates:
        if Path(path).exists(): return ImageFont.truetype(path, size)
    return ImageFont.load_default()

def arrow(draw, start, end, color=(70, 70, 84), width=3):
    draw.line([start, end], fill=color, width=width)
    x1, y1 = end; x0, y0 = start
    if abs(x1-x0) >= abs(y1-y0):
        s = 1 if x1 > x0 else -1
        draw.polygon([(x1,y1),(x1-12*s,y1-7),(x1-12*s,y1+7)], fill=color)
    else:
        s = 1 if y1 > y0 else -1
        draw.polygon([(x1,y1),(x1-7,y1-12*s),(x1+7,y1-12*s)], fill=color)

def box(draw, xy, title, lines, fill, outline=(49,45,87)):
    x1,y1,x2,y2 = xy
    draw.rounded_rectangle(xy, radius=16, fill=fill, outline=outline, width=3)
    tf = font(25, True); bf = font(18)
    draw.text((x1+18,y1+18), title, font=tf, fill=(25,29,55))
    yy=y1+55
    for line in lines:
        draw.text((x1+18,yy), line, font=bf, fill=(55,58,78)); yy += 25

def make_architecture():
    im = Image.new('RGB', (1500, 720), 'white'); d = ImageDraw.Draw(im)
    d.text((55,30), 'Arkitektura e përgjithshme e DataGuard', font=font(32, True), fill=(25,29,55))
    box(d,(70,230,340,440),'Përdoruesi',['Admin ose User','Shfletues web'],(236,245,241))
    box(d,(470,175,840,495),'Frontend',['React + TypeScript','Vite','Tailwind CSS','Komponentë dhe faqe'],(238,234,250))
    box(d,(990,105,1400,315),'Supabase',['Auth dhe JWT','REST API / supabase-js','Kontroll i sesionit'],(225,248,238))
    box(d,(990,405,1400,615),'PostgreSQL',['Tabelat e modulit','RLS policies','Marrëdhënie dhe ruajtje'],(235,241,250))
    arrow(d,(340,335),(470,335)); arrow(d,(840,285),(990,210)); arrow(d,(840,390),(990,510)); arrow(d,(1195,315),(1195,405))
    d.text((365,300),'HTTPS',font=font(17),fill=(80,80,95)); d.text((855,235),'Kyçje / token',font=font(17),fill=(80,80,95)); d.text((855,470),'Lexim / shkrim',font=font(17),fill=(80,80,95)); d.text((1218,345),'RLS',font=font(17),fill=(80,80,95))
    im.save(ASSETS/'architecture.png')

def entity(draw, xy, title, fields):
    x1,y1,x2,y2=xy
    draw.rounded_rectangle(xy, radius=10, fill=(249,249,252), outline=(49,45,87), width=2)
    draw.rounded_rectangle((x1,y1,x2,y1+34),radius=10,fill=(82,232,173),outline=(49,45,87),width=2)
    draw.rectangle((x1,y1+24,x2,y1+34),fill=(82,232,173))
    draw.text((x1+12,y1+7),title,font=font(16,True),fill=(25,29,55))
    yy=y1+43
    for f in fields:
        draw.text((x1+12,yy),f,font=font(13),fill=(55,58,78)); yy+=18

def make_er():
    im=Image.new('RGB',(1600,1070),'white'); d=ImageDraw.Draw(im)
    d.text((45,25),'Modeli konceptual i bazës së të dhënave',font=font(30,True),fill=(25,29,55))
    entity(d,(70,115,335,255),'profiles',['id (PK)','full_name','role','created_at'])
    entity(d,(520,105,830,255),'privacy_policies',['id (PK)','title','status','created_by (FK)'])
    entity(d,(1080,105,1410,255),'policy_versions',['id (PK)','policy_id (FK)','version','effective_date'])
    entity(d,(570,430,860,590),'consents',['id (PK)','user_id (FK)','policy_version_id (FK)','status, decided_at'])
    entity(d,(80,430,340,570),'audit_logs',['id (PK)','actor_id (FK)','action, resource','status, created_at'])
    entity(d,(1085,430,1410,590),'data_requests',['id (PK)','user_id (FK)','type, status','reviewed_by (FK)'])
    entity(d,(80,755,345,900),'retention_policies',['id (PK)','data_type','retention_period','action, status'])
    entity(d,(600,750,860,900),'customers',['id (PK)','name, email','phone, address','created_at'])
    arrow(d,(335,180),(520,180)); arrow(d,(830,180),(1080,180)); arrow(d,(205,255),(205,430)); arrow(d,(335,225),(570,470)); arrow(d,(1245,255),(850,470))
    d.text((385,145),'krijon',font=font(13),fill=(80,80,95)); d.text((930,145),'ka versione',font=font(13),fill=(80,80,95)); d.text((375,340),'jep miratim',font=font(13),fill=(80,80,95)); d.text((955,340),'lidhet me versionin',font=font(13),fill=(80,80,95)); d.text((220,330),'kryen veprime',font=font(13),fill=(80,80,95))
    d.text((80,950),'Marrëdhëniet kryesore: një politikë ka shumë versione; një version mund të ketë shumë miratime;',font=font(16),fill=(55,58,78))
    d.text((80,978),'ndërsa një profil mund të krijojë politika, miratime, audit logs dhe kërkesa për të dhëna.',font=font(16),fill=(55,58,78))
    im.save(ASSETS/'database_model.png')

make_architecture(); make_er()

doc=Document(); sec=doc.sections[0]
sec.page_width,sec.page_height=Cm(21),Cm(29.7)
sec.top_margin,sec.bottom_margin=Cm(2.5),Cm(2.5)
sec.left_margin,sec.right_margin=Cm(3),Cm(3)
sec.header_distance,sec.footer_distance=Cm(1.25),Cm(1.25)
normal=doc.styles['Normal']; normal.font.name='Times New Roman'; normal._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); normal._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman'); normal._element.rPr.rFonts.set(qn('w:cs'),'Times New Roman'); normal.font.size=Pt(12); normal.paragraph_format.line_spacing=1.5; normal.paragraph_format.space_after=Pt(0); normal.paragraph_format.first_line_indent=Cm(1.25); normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for name,size in [('Heading 1',14),('Heading 2',12)]:
    s=doc.styles[name]; s.font.name='Times New Roman'; s._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); s._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman'); s.font.size=Pt(size); s.font.bold=True; s.paragraph_format.line_spacing=1.5; s.paragraph_format.space_before=Pt(12); s.paragraph_format.space_after=Pt(6); s.paragraph_format.first_line_indent=Cm(0); s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
def heading(t,l): doc.add_paragraph(t,style=f'Heading {l}')
def para(t):
    p=doc.add_paragraph(style='Normal'); p.add_run(t); return p
def caption(t):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent=Cm(0); p.paragraph_format.space_after=Pt(6); r=p.add_run(t); r.font.name='Times New Roman'; r._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); r.font.size=Pt(11); r.font.italic=True
def set_cell(cell,text,bold=False):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.first_line_indent=Cm(0); p.paragraph_format.line_spacing=1.15; p.paragraph_format.space_after=Pt(0); r=p.add_run(text); r.font.name='Times New Roman'; r._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); r.font.size=Pt(10); r.font.bold=bold; cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
def table_geometry(table,widths):
    table.alignment=WD_TABLE_ALIGNMENT.LEFT; table.autofit=False; tp=table._tbl.tblPr; tw=tp.first_child_found_in('w:tblW'); tw.set(qn('w:w'),str(sum(widths))); tw.set(qn('w:type'),'dxa'); ind=tp.first_child_found_in('w:tblInd')
    if ind is None: ind=OxmlElement('w:tblInd'); tp.append(ind)
    ind.set(qn('w:w'),'120'); ind.set(qn('w:type'),'dxa')
    for g,w in zip(table._tbl.tblGrid.gridCol_lst,widths): g.set(qn('w:w'),str(w))
    for row in table.rows:
        for cell,w in zip(row.cells,widths): cell._tc.tcPr.tcW.set(qn('w:w'),str(w)); cell._tc.tcPr.tcW.set(qn('w:type'),'dxa')

heading('4. ARKITEKTURA DHE DIZAJNI I ZGJIDHJES',1)

heading('4.1 Arkitektura e sistemit',2)
para('DataGuard është projektuar me arkitekturë web klient-server, ku ndërfaqja e përdoruesit ekzekutohet në shfletues, ndërsa autentifikimi, API-ja dhe ruajtja e të dhënave ofrohen nga Supabase. Kjo qasje e ndan qartë përgjegjësinë e paraqitjes së të dhënave nga përgjegjësia e ruajtjes dhe kontrollit të tyre. Aplikacioni frontend komunikon me Supabase përmes bibliotekës supabase-js dhe përdor kërkesa HTTPS për funksionet e autentifikimit dhe operacionet mbi të dhënat.')
para('Përdoruesi ose administratori e hap aplikacionin nga shfletuesi. React e paraqet ndërfaqen si komponentë dhe faqe, ndërsa TypeScript ndihmon në kontrollin e tipeve të të dhënave gjatë zhvillimit. Supabase Auth verifikon përdoruesin dhe krijon sesionin e tij. Kërkesat për lexim ose shkrim drejtohen te PostgreSQL, ku rregullat Row Level Security kontrollojnë nëse përdoruesi ka të drejtë të qaset në rreshtat përkatës. Kjo strukturë paraqitet në Figurën 1.')
doc.add_picture(str(ASSETS/'architecture.png'),width=Inches(6.3)); caption('Figura 1. Arkitektura e përgjithshme e sistemit DataGuard (burimi: përpunim i autores).')
para('Zgjedhja e Supabase si backend e bën arkitekturën të përshtatshme për një prototip Bachelor. Nuk kërkohet ndërtimi i një serveri të veçantë me Laravel ose .NET, por funksionet e nevojshme backend realizohen përmes Auth, PostgreSQL, API-së së gjeneruar dhe RLS. Supabase Auth përdor JSON Web Tokens për autentifikim dhe lidhet me mekanizmat e databazës për autorizim [4]. Kjo zgjidhje ruan thjeshtësinë e projektit, ndërsa demonstron një arkitekturë moderne të aplikacionit web.')

heading('4.2 Teknologjitë e përdorura',2)
para('Teknologjitë janë zgjedhur duke u bazuar në kërkesat e projektit, thjeshtësinë e implementimit dhe mundësinë e ndërtimit të një ndërfaqeje profesionale. React përdoret për ndërtimin e ndërfaqes së përdoruesit. Ai lejon organizimin e ekranit në komponentë të ripërdorshëm, ku secili komponent mund të ketë logjikën dhe pamjen e vet [1]. Kjo është e përshtatshme për DataGuard, sepse shumë pjesë të ndërfaqes, si kartelat, tabelat, badges dhe butonat, përdoren në disa faqe.')
para('TypeScript përdoret si zgjerim i JavaScript-it me sistem tipeve. Ai ndihmon në evidentimin e mospërputhjeve të të dhënave gjatë zhvillimit dhe ul mundësinë e gabimeve në kod [2]. Vite përdoret si mjedis zhvillimi dhe mjet ndërtimi. Dokumentacioni zyrtar e përshkruan Vite si mjet që ofron server të shpejtë zhvillimi dhe proces ndërtimi për asetet e aplikacionit [3].')
para('Tailwind CSS përdoret për stilizimin e ndërfaqes dhe përshtatjen e saj në madhësi të ndryshme të ekranit. Ai mbështetet në utility classes dhe variante responsive, të cilat mundësojnë krijimin e ndërfaqeve adaptive [6]. Në projekt janë përdorur tone vjollcë të errët, vjollcë e hapur dhe jeshile mint për të krijuar identitetin vizual të DataGuard.')
para('Supabase ofron funksionet backend të nevojshme për projektin: autentifikimin e përdoruesve, ruajtjen e të dhënave në PostgreSQL, API për qasje në tabela dhe politikat RLS. PostgreSQL është sistem i menaxhimit të bazës së të dhënave objekt-relacionale, me mbështetje për SQL, foreign keys, triggers dhe integritet transaksional [5]. Këto karakteristika janë të përshtatshme për tabelat e lidhura të DataGuard, sidomos për versionet e politikave dhe miratimet.')

caption('Tabela 2. Teknologjitë e përdorura në DataGuard')
tbl=doc.add_table(rows=1,cols=3); tbl.style='Table Grid'
for cell,text in zip(tbl.rows[0].cells,['Teknologjia','Roli në projekt','Arsyeja e përdorimit']): set_cell(cell,text,True)
for row in [
    ('React','Ndërfaqja e përdoruesit','Komponentë të ripërdorshëm dhe faqe interaktive.'),
    ('TypeScript','Kontrolli i tipeve','Rrit qartësinë dhe sigurinë e kodit.'),
    ('Vite','Mjedisi i zhvillimit','Nisje e shpejtë dhe build për aplikacion web.'),
    ('Tailwind CSS','Stilizimi','Ndërfaqe responsive dhe konsistente.'),
    ('Supabase','Backend i menaxhuar','Auth, API dhe lidhje me databazën.'),
    ('PostgreSQL','Baza e të dhënave','Ruajtje relacionale e entiteteve të modulit.')]:
    cells=tbl.add_row().cells
    for c,t in zip(cells,row): set_cell(c,t)
table_geometry(tbl,[1750,3050,4560])

heading('4.3 Arkitektura e aplikacionit',2)
para('Në nivel aplikacioni, DataGuard organizohet në disa shtresa të thjeshta. Shtresa e paraqitjes përmban layout-in e përgjithshëm, sidebar-in, topbar-in, kartelat, formularët, tabelat dhe komponentët e njoftimeve. Shtresa e logjikës së ndërfaqes menaxhon gjendjen e faqes së hapur, formularët, filtrat, mesazhet dhe reagimin ndaj veprimeve të përdoruesit. Shtresa e qasjes në të dhëna lidhet me klientin Supabase dhe kryen operacionet select, insert, update dhe delete.')
para('Për të ruajtur përvojë të qartë, menuja anësore paraqet faqe të ndryshme sipas rolit. Administratori sheh Dashboard, Privacy Policies, Consents, Audit Logs, Retention Policies, Data Requests, My Privacy dhe SME Customer App. Përdoruesi i zakonshëm sheh vetëm Privacy Policies dhe My Privacy. Ky kontroll vizual nuk është i vetmi mekanizëm sigurie; ai plotësohet nga rregullat e autorizimit në databazë.')
para('Faqja Dashboard paraqet statistika të përgjithshme dhe aktivitetet e fundit. Privacy Policies përmban krijimin, shikimin, ndryshimin dhe versionimin e politikave. Consents paraqet vendimet e përdoruesve, ndërsa Audit Logs paraqet historinë e veprimeve. Retention Policies dhe Data Requests mbulojnë periudhën e ruajtjes së të dhënave dhe procesimin e kërkesave. SME Customer App shërben si shembull konkret i aplikacionit që përdor modulin e privatësisë.')

heading('4.4 Dizajni i bazës së të dhënave',2)
para('Dizajni i databazës është ndërtuar duke filluar nga kërkesat funksionale. Tabela profiles ruan emrin, rolin dhe identifikuesin e përdoruesit. Ky identifikues lidhet me përdoruesin përkatës në Supabase Auth. Tabela privacy_policies ruan dokumentet kryesore të privatësisë, ndërsa tabela policy_versions ruan versionet dhe tekstin e secilit version. Kjo ndarje e bën të mundur ruajtjen e historikut pa e humbur përmbajtjen e versioneve të vjetra.')
para('Tabela consents lidh përdoruesin me policy_versions dhe ruan statusin Accepted, Rejected ose Withdrawn, si dhe datën e vendimit. Tabela audit_logs ruan aktorin, veprimin, burimin, statusin dhe kohën e krijimit. Tabela retention_policies definon kategorinë e të dhënave, periudhën e ruajtjes, veprimin dhe statusin. Data_requests ruan kërkesat për eksport ose fshirje, ndërsa customers ruan të dhënat e klientëve të aplikacionit demonstrues.')
para('Përdorimi i primary keys dhe foreign keys mbron lidhjet kryesore ndërmjet entiteteve. Për shembull, çdo version i politikës lidhet me një politikë ekzistuese, ndërsa çdo miratim lidhet me një përdorues dhe me një version të caktuar. Trigger-i handle_new_user krijon profil bazë kur regjistrohet përdorues i ri në Supabase Auth. Kjo shmang nevojën që profili të krijohet manualisht pas çdo regjistrimi.')
doc.add_picture(str(ASSETS/'database_model.png'),width=Inches(6.3)); caption('Figura 2. Modeli konceptual i bazës së të dhënave (burimi: përpunim i autores).')

heading('4.5 Marrëdhëniet ndërmjet entiteteve',2)
para('Marrëdhënia qendrore e modelit është lidhja ndërmjet privacy_policies dhe policy_versions. Një politikë mund të ketë një ose më shumë versione, ndërsa një version i përket vetëm një politike. Kjo është marrëdhënie një-me-shumë. Në të njëjtën mënyrë, një policy_version mund të lidhet me shumë consents, sepse shumë përdorues mund të marrin vendim për të njëjtin version.')
para('Një profil mund të ketë shumë miratime dhe shumë data requests. Profili mund të paraqitet gjithashtu si krijues i politikës ose si aktor në audit log. Për kërkesat e të dhënave, fusha reviewed_by lejon ruajtjen e administratorit që e ka shqyrtuar kërkesën. Tabelat retention_policies dhe customers janë të pavarura nga përdoruesi specifik në këtë prototip, por lidhen me konceptin e përgjithshëm të administrimit të të dhënave të biznesit.')
para('Ky model mbështet zgjerim të ardhshëm. Për shembull, mund të shtohen tabela për kategori të avancuara të të dhënave, njoftime, dokumente të eksportuara ose histori të ndryshimeve të klientëve. Megjithatë, modeli aktual ruan vetëm entitetet e domosdoshme për funksionet e kërkuara në temën e diplomës.')

heading('4.6 Dizajni i ndërfaqes',2)
para('Dizajni i ndërfaqes ndjek parimin e qartësisë dhe reduktimit të kompleksitetit. Faqet janë ndërtuar me sfond të hapur, karta të bardha, sidebar vjollcë të errët dhe thekse vjollcë të hapur e jeshile mint. Sidebar-i mban navigimin kryesor dhe informacionin e llogarisë, ndërsa zona qendrore paraqet përmbajtjen e modulit. Kjo strukturë përsëritet në të gjitha faqet dhe e bën aplikacionin të parashikueshëm për përdoruesin.')
para('Tabelat përdoren për regjistrimet që kërkojnë krahasim, si politikat, miratimet, audit logs, kërkesat për të dhëna dhe klientët. Formularët paraqiten vetëm kur përdoruesi zgjedh të krijojë ose të ndryshojë të dhëna. Butonat janë të dallueshëm sipas veprimit: veprimet kryesore shfaqen me ngjyrë mint, ndërsa veprimet dytësore me stil më të qetë. Mesazhet e njoftimit konfirmojnë suksesin e veprimeve.')
para('Për përdoruesin e zakonshëm, faqja My Privacy e grupon informacionin në tri pjesë: politika aktuale, statusi i miratimit dhe kërkesat për të dhëna. Kjo e redukton numrin e hapave dhe e bën të qartë se cilat veprime mund të kryejë përdoruesi. Dizajni responsive siguron që përmbajtja të mbetet e lexueshme edhe në ekrane më të vogla.')

heading('4.7 Autentifikimi dhe kontrolli i qasjes',2)
para('Autentifikimi në DataGuard realizohet me email dhe fjalëkalim përmes Supabase Auth. Pas krijimit të llogarisë, trigger-i në databazë krijon automatikisht një rresht në tabelën profiles me rolin fillestar user. Roli admin caktohet për përdoruesit që duhet të menaxhojnë pjesët administrative të sistemit. Kjo qasje e ndan identitetin e përdoruesit, i cili ruhet në Auth, nga informacioni shtesë i profilit dhe rolit, i cili ruhet në public.profiles.')
para('Autorizimi zbatohet në dy nivele. Niveli i parë është ndërfaqja: navigimi dhe faqet administrative nuk shfaqen për përdoruesit e zakonshëm. Niveli i dytë është databaza: tabelat kanë Row Level Security të aktivizuar dhe politikat SQL përdorin identitetin auth.uid() dhe funksionin is_admin(). RLS ofron rregulla autorizimi që ekzekutohen brenda PostgreSQL dhe mund të kombinohen me Supabase Auth për mbrojtjen e qasjes nga shfletuesi deri te databaza [7].')
para('Për shembull, përdoruesi mund të menaxhojë vetëm consents dhe data_requests që lidhen me identitetin e tij, ndërsa administratori mund të menaxhojë politikat, versionet, retention policies dhe customers. Dokumentet e politikave mund të lexohen nga të gjithë përdoruesit e autentifikuar, sepse përdoruesi duhet të ketë mundësi ta lexojë politikën para se të marrë vendim. Ky model është i mjaftueshëm për prototipin dhe demonstron parimin e kontrollit të qasjes me role.')

heading('Referenca të përdorura në këtë kapitull',2)
refs=[
    '[1] React. React Documentation: Components and User Interfaces. Në dispozicion: https://react.dev/',
    '[2] Microsoft. TypeScript Documentation: TypeScript for JavaScript Programmers. Në dispozicion: https://www.typescriptlang.org/docs/handbook/typescript-in-5-minutes.html',
    '[3] Vite. Getting Started. Në dispozicion: https://vite.dev/guide/',
    '[4] Supabase. Auth Documentation. Në dispozicion: https://supabase.com/docs/guides/auth',
    '[5] PostgreSQL Global Development Group. PostgreSQL Documentation: What is PostgreSQL? Në dispozicion: https://www.postgresql.org/docs/current/intro-whatis.html',
    '[6] Tailwind CSS. Responsive Design. Në dispozicion: https://tailwindcss.com/docs/responsive-design',
    '[7] Supabase. Row Level Security Documentation. Në dispozicion: https://supabase.com/docs/guides/database/postgres/row-level-security'
]
for t in refs:
    p=doc.add_paragraph(style='Normal'); p.paragraph_format.first_line_indent=Cm(0); p.paragraph_format.space_after=Pt(3); p.add_run(t)
foot=sec.footer.paragraphs[0]; foot.alignment=WD_ALIGN_PARAGRAPH.CENTER; rr=foot.add_run('DataGuard - Punim diplome'); rr.font.name='Times New Roman'; rr._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); rr.font.size=Pt(10)
OUT.parent.mkdir(exist_ok=True); doc.save(OUT); print(OUT)
