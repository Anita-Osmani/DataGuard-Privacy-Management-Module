from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUT = Path('outputs/Kapitulli_5_Miratimet_dhe_Tekstet_e_Privatesise.docx')
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.5)
sec.left_margin, sec.right_margin = Cm(3), Cm(3)
sec.header_distance, sec.footer_distance = Cm(1.25), Cm(1.25)

normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
normal._element.rPr.rFonts.set(qn('w:cs'), 'Times New Roman')
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.first_line_indent = Cm(1.25)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for style_name, size in [('Heading 1', 14), ('Heading 2', 12)]:
    s = doc.styles[style_name]
    s.font.name = 'Times New Roman'
    s._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
    s._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    s.font.size = Pt(size)
    s.font.bold = True
    s.paragraph_format.line_spacing = 1.5
    s.paragraph_format.space_before = Pt(12)
    s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.first_line_indent = Cm(0)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

def heading(text, level):
    doc.add_paragraph(text, style=f'Heading {level}')

def para(text):
    p = doc.add_paragraph(style='Normal')
    p.add_run(text)
    return p

heading('5. MODULI I MIRATIMEVE DHE TEKSTEVE TË PRIVATËSISË', 1)

heading('5.1 Përshkrimi i modulit', 2)
para('Moduli i miratimeve dhe teksteve të privatësisë paraqet funksionin kryesor të DataGuard, sepse lidh dokumentin e privatësisë me vendimin e përdoruesit. Në vend që politika e privatësisë të paraqitet vetëm si tekst statik në aplikacion, moduli e trajton atë si dokument të menaxhueshëm, me version, status dhe datë të hyrjes në fuqi. Kjo i jep administratorit mundësi të organizojë tekstet e privatësisë, ndërsa përdoruesit i jep informacion të qartë për politikën që po pranon.')
para('Përmes këtij moduli administratori mund të krijojë politikë të re, të shikojë politikën ekzistuese, të ndryshojë titullin dhe të krijojë version të ri. Përdoruesi e sheh politikën aktuale në faqen My Privacy dhe mund të japë ose të tërheqë miratimin. Sistemi ruan lidhjen e vendimit me versionin konkret të dokumentit. Në këtë mënyrë, një vendim nuk lidhet vetëm me emrin e politikës, por me përmbajtjen dhe versionin që ishte i vlefshëm në kohën e marrjes së vendimit.')
para('Moduli është projektuar për një NVM që dëshiron të administrojë privatësinë pa përdorur procese të ndërlikuara. Ai nuk synon të zëvendësojë këshillimin juridik ose një platformë enterprise për pajtueshmëri. Qëllimi është të demonstrojë një rrjedhë të kuptueshme: administratori publikon politikën, përdoruesi e lexon dhe merr vendim, ndërsa sistemi ruan historikun për qëllime administrimi dhe auditimi.')

heading('5.2 Menaxhimi i politikave të privatësisë', 2)
para('Faqja Privacy Policies përdoret nga administratori për menaxhimin e dokumenteve të privatësisë. Lista e politikave paraqet titullin, versionin aktual, statusin, datën e hyrjes në fuqi dhe datën e përditësimit. Kjo paraqitje e bën të mundur që administratori të ketë pasqyrë të shpejtë mbi dokumentet aktive dhe dokumentet që janë ende në përgatitje.')
para('Gjatë krijimit të politikës, administratori plotëson titullin, versionin fillestar, datën e hyrjes në fuqi dhe tekstin e politikës. Këto të dhëna ruhen në tabelat privacy_policies dhe policy_versions. Tabela e parë ruan identitetin e dokumentit, ndërsa tabela e dytë ruan tekstin dhe informacionin që lidhet me versionin. Ndarja e këtyre tabelave është e nevojshme që një politikë të ketë disa versione pa humbur identitetin e saj kryesor.')
para('Statuset Draft, Current dhe Archived përdoren për të treguar gjendjen e politikës. Draft shërben për dokument që ende nuk duhet të paraqitet si politikë aktive. Current përdoret për politikën që i shfaqet përdoruesit në My Privacy. Archived ruan dokumentin që nuk përdoret më si aktual, por që mund të jetë i nevojshëm për historik. Në prototip, përdorimi i këtyre statuseve e bën administrimin e dokumenteve më të qartë, edhe pse rrjedha është mbajtur e thjeshtë.')
para('Funksioni View lejon kontrollimin e tekstit të ruajtur, ndërsa Edit përdoret për ndryshime të kufizuara të politikës. Krijimi i versionit të ri bëhet përmes butonit New Version. Ky ndalim i ndryshimit të drejtpërdrejtë të përmbajtjes së versionit të vjetër ndihmon që historia e dokumentit të mos humbet. Në punë reale, teksti i politikës duhet të përgatitet ose të rishikohet nga person i autorizuar; DataGuard demonstron vetëm administrimin teknik të këtij teksti.')

heading('5.3 Versionimi i politikave', 2)
para('Versionimi është procesi i ruajtjes së ndryshimeve në një dokument me anë të numrave të versioneve. Në DataGuard, administratori mund të krijojë versionin 1.0 gjatë publikimit të parë të politikës dhe, më vonë, versionin 1.1 ose 2.0 kur përmbajtja duhet të përditësohet. Versionet e mëparshme nuk fshihen nga sistemi; ato mbeten pjesë e historikut të politikës.')
para('Për secilin version ruhen numri i versionit, teksti, data e hyrjes në fuqi dhe data e krijimit. Çdo version lidhet me një politikë përmes policy_id. Për më tepër, kufizimi i unikësisë për kombinimin policy_id dhe version siguron që brenda së njëjtës politikë të mos krijohen dy versione me emër të njëjtë. Kjo parandalon paqartësinë në lidhjen ndërmjet politikës dhe miratimeve të përdoruesve.')
para('Për shembull, nëse politika fillestare “Customer Privacy Policy” ndryshohet për të sqaruar periudhën e ruajtjes së të dhënave, administratori krijon version të ri në vend që ta zëvendësojë tekstin e vjetër. Përdoruesi që e ka pranuar versionin 1.0 vazhdon të ketë regjistrim të lidhur me versionin 1.0. Përdoruesi që pranon dokumentin pas ndryshimit lidhet me versionin e ri. Kjo është pjesë e rëndësishme e gjurmueshmërisë së miratimeve.')
para('Versionimi i politikave nuk do të thotë që sistemi detyron automatikisht çdo përdorues të pranojë dokument të ri. Në prototip, qëllimi është ruajtja e historisë dhe paraqitja e versionit aktual. Njoftimet e avancuara për version të ri, miratimet e detyrueshme sipas llojit të ndryshimit ose proceset juridike të rivlerësimit mbeten si mundësi për zgjerime të ardhshme.')

heading('5.4 Menaxhimi i miratimeve', 2)
para('Miratimi në DataGuard regjistrohet në tabelën consents. Çdo rresht përmban identifikuesin e përdoruesit, identifikuesin e versionit të politikës, statusin dhe datën e vendimit. Statuset e mbështetura janë Accepted, Rejected dhe Withdrawn. Ky model është i mjaftueshëm për të paraqitur vendimin kryesor të përdoruesit lidhur me politikën e privatësisë.')
para('Në faqen My Privacy përdoruesi sheh titullin dhe versionin e politikës aktuale. Nëse nuk ekziston miratim aktiv, shfaqet butoni Accept Privacy Policy. Pas pranimit, ndërfaqja paraqet statusin Consent active dhe tregon versionin e politikës së pranuar. Nëse përdoruesi vendos të mos vazhdojë me miratimin, ai mund të zgjedhë Withdraw Consent. Kjo e bën statusin e miratimit të dukshëm dhe të kontrollueshëm nga vetë përdoruesi.')
para('Në panelin administrativ, faqja Consents paraqet listën e miratimeve me kolonat User, Policy, Version, Status dhe Date. Kjo i ndihmon administratorit të verifikojë nëse një përdorues e ka pranuar politikën aktuale dhe cilin version e ka pranuar. Lista nuk shfaq të dhëna të panevojshme personale; ajo përqendrohet te informacioni që nevojitet për administrimin e vendimeve të privatësisë.')
para('Udhëzimet e EDPB-së për consent theksojnë se miratimi, kur përdoret si bazë për përpunim, duhet të jetë i vlefshëm dhe i demonstrueshëm [2]. DataGuard e mbështet këtë kërkesë në nivel të prototipit duke ruajtur personin, versionin e politikës, statusin dhe kohën e vendimit. Sistemi nuk vlerëson automatikisht nëse miratimi është baza e duhur ligjore për çdo aktivitet të biznesit; ky vendim kërkon analizë të kontekstit konkret.')

heading('5.5 Historiku dhe tërheqja e miratimeve', 2)
para('Historiku i miratimeve e bën të mundur që ndryshimet e vendimit të përdoruesit të mos humben. Nëse përdoruesi pranon politikën dhe më vonë e tërheq miratimin, sistemi regjistron statusin Withdrawn. Kjo paraqet gjendjen e re të vendimit, ndërsa regjistrimi përmban edhe datën kur është kryer veprimi. Administratori mund ta shohë këtë informacion në faqen Consents dhe në Audit Logs kur veprimi është regjistruar si aktivitet i rëndësishëm.')
para('Tërheqja e miratimit nuk duhet të jetë më e vështirë sesa pranimi i tij. Për këtë arsye, në DataGuard butoni Withdraw Consent vendoset në të njëjtën faqe ku përdoruesi e sheh statusin e consent-it. Kjo rrjedhë është e qëllimshme: përdoruesi nuk duhet të kërkojë në shumë faqe ose të kontaktojë administratorin vetëm për ta ndryshuar vendimin e tij. Udhëzimet e EDPB-së e trajtojnë lehtësinë e tërheqjes si element të rëndësishëm të menaxhimit të miratimit [2].')
para('Tërheqja e miratimit nuk nënkupton automatikisht fshirje të llogarisë ose fshirje të të gjitha të dhënave. Në aplikacion këto funksione janë të ndara: miratimi menaxhohet në My Privacy, ndërsa kërkesa për fshirje krijohet përmes funksionit Request Data Deletion. Kjo ndarje ndihmon përdoruesin të kuptojë se ndryshimi i consent-it dhe kërkesa për fshirje janë veprime të ndryshme, të cilat mund të kenë rrjedha administrative të ndryshme.')

heading('5.6 Implementimi dhe rezultatet', 2)
para('Moduli është implementuar me React në anën e ndërfaqes dhe me Supabase/PostgreSQL në anën e ruajtjes së të dhënave. Formulari Create Policy krijon një rresht në privacy_policies dhe versionin e parë në policy_versions. Formulari New Version krijon rresht të ri në policy_versions dhe e ruan versionin e mëparshëm. Kjo logjikë u testua duke krijuar politika dhe versione të reja në panelin e administratorit.')
para('Veprimet e përdoruesit për pranimin ose tërheqjen e consent-it shkruhen në tabelën consents, së bashku me versionin e politikës. Pas ruajtjes, My Privacy përditëson statusin që i shfaqet përdoruesit. Në anën administrative, regjistrimet lexohen dhe paraqiten në tabelën e miratimeve. Ky rezultat demonstron se ekziston lidhje e ruajtur ndërmjet tekstit të politikës dhe vendimit të përdoruesit.')
para('Nga testimi manual u vërtetua se përdoruesi mund të shikojë politikën aktive, të pranojë versionin aktual dhe të tërheqë miratimin. U vërtetua gjithashtu se administratori mund të krijojë politikë dhe version të ri pa fshirë historinë e versioneve të mëparshme. Këto rezultate përmbushin qëllimin e modulit: ofrimin e një mekanizmi të thjeshtë, të kuptueshëm dhe të gjurmueshëm për administrimin e politikave dhe miratimeve në një aplikacion SME.')

heading('Referenca të përdorura në këtë kapitull', 2)
for reference in [
    '[1] European Parliament and Council of the European Union. Regulation (EU) 2016/679 (General Data Protection Regulation), 27 April 2016. EUR-Lex. Në dispozicion: https://eur-lex.europa.eu/eli/reg/2016/679/oj',
    '[2] European Data Protection Board. Guidelines 05/2020 on consent under Regulation 2016/679, Version 1.1, 4 May 2020. Në dispozicion: https://www.edpb.europa.eu/our-work-tools/our-documents/guidelines/guidelines-052020-consent-under-regulation-2016679_en'
]:
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(3)
    p.add_run(reference)

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run('DataGuard - Punim diplome')
run.font.name = 'Times New Roman'
run._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
run.font.size = Pt(10)
OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(OUT)
